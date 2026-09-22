"""Professional single-column ATS-friendly DOCX resume renderer."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from packages.shared.schemas import ParsedResume

INK = RGBColor(25, 31, 45)
MUTED = RGBColor(70, 76, 88)


def _set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = INK


def _add_section_heading(doc: Document, title: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(9)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    _set_run_font(paragraph.add_run(title.upper()), 10.5, True)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "315C6D")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _add_body(doc: Document, text: str, *, bullet: bool = False) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if bullet else None)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.keep_together = True
    if bullet:
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    _set_run_font(paragraph.add_run(text))


def render_ats_resume(data: ParsedResume, output_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    style = doc.styles["Normal"]
    style.font.name = "Aptos"
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
    style.font.size = Pt(10.5)
    style.font.color.rgb = INK

    if data.name:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        _set_run_font(paragraph.add_run(data.name.upper()), 17, True)

    contact_bits = [
        data.location,
        data.contact.email,
        data.contact.phone,
        data.contact.linkedin,
        data.contact.github,
        data.contact.portfolio,
    ]
    contact_line = "  |  ".join(bit for bit in contact_bits if bit)
    if contact_line:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(contact_line)
        _set_run_font(run, 9)
        run.font.color.rgb = MUTED

    if data.summary:
        _add_section_heading(doc, "Professional Summary")
        _add_body(doc, data.summary)

    all_skills = list(dict.fromkeys(data.skills_must_have + data.skills_nice_to_have))
    if all_skills:
        _add_section_heading(doc, "Technical Skills")
        _add_body(doc, " • ".join(all_skills))

    if data.experience:
        _add_section_heading(doc, "Professional Experience")
        for experience in data.experience:
            dates = " - ".join(
                value for value in [experience.start_date, experience.end_date] if value
            )
            header = " | ".join(
                value for value in [experience.role, experience.company] if value
            )
            if dates:
                header = f"{header}  ({dates})" if header else dates
            if header:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.keep_with_next = True
                _set_run_font(paragraph.add_run(header), 10.5, True)
            for bullet in experience.bullets:
                _add_body(doc, bullet, bullet=True)

    if data.projects:
        _add_section_heading(doc, "Selected Projects")
        for project in data.projects:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.keep_with_next = True
            heading = project.name
            if project.technologies:
                heading += f" | {', '.join(project.technologies)}"
            _set_run_font(paragraph.add_run(heading), 10.5, True)
            if project.description:
                _add_body(doc, project.description)
            for bullet in project.bullets:
                _add_body(doc, bullet, bullet=True)

    if data.education:
        _add_section_heading(doc, "Education")
        for education in data.education:
            line = " | ".join(
                value
                for value in [education.degree, education.institution, education.year]
                if value
            )
            if line:
                _add_body(doc, line)

    if data.certifications:
        _add_section_heading(doc, "Certifications")
        for certification in data.certifications:
            _add_body(doc, certification, bullet=True)

    if data.achievements:
        _add_section_heading(doc, "Achievements")
        for achievement in data.achievements:
            _add_body(doc, achievement, bullet=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
