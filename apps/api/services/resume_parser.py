"""Parse PDF/DOCX resumes into structured JSON."""
import re
from pathlib import Path

import pdfplumber
from docx import Document

from packages.shared.schemas import (
    ContactInfo,
    EducationEntry,
    ExperienceEntry,
    ParsedResume,
    ProjectEntry,
)

SECTION_KEYWORDS: dict[str, tuple[str, ...]] = {
    "experience": ("WORKEXPERIENCE", "PROFESSIONALEXPERIENCE", "EXPERIENCE", "EMPLOYMENT", "CAREERHISTORY", "WORKHISTORY", "INTERNSHIPS"),
    "skills": ("TECHNICALSKILLS", "CORESKILLS", "SKILLS", "TECHNOLOGIES", "TECHSTACK", "TECHNICALPROFICIENCIES"),
    "education": ("EDUCATION", "ACADEMIC", "QUALIFICATIONS"),
    "certifications": ("CERTIFICATIONS", "CERTIFICATION", "LICENSES"),
    "projects": ("PROJECTS", "PERSONALPROJECTS", "OPENSOURCE"),
    "achievements": ("ACHIEVEMENTS", "AWARDS", "AWARDSANDACHIEVEMENTS"),
    "summary": ("SUMMARY", "PROFESSIONALSUMMARY", "OBJECTIVE", "PROFILE", "ABOUTME"),
}

BULLET_PREFIX = re.compile(r"^\s*[-•*▪◦‣·]\s*|^\s*\d+[.)]\s*")
MONTH = r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
DATE_VALUE = rf"(?:{MONTH}\s+)?(?:19|20)\d{{2}}"
DATE_RANGE = re.compile(
    rf"(?P<start>{DATE_VALUE})\s*(?:-|–|—|to)\s*"
    rf"(?P<end>{DATE_VALUE}|present|current)",
    re.IGNORECASE,
)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(\+?\d[\d\s().-]{8,}\d)")
URL_RE = re.compile(
    r"(https?://[^\s|,;]+|(?:www\.|linkedin\.com/|github\.com/)[^\s|,;]+|"
    r"(?<![@\w])(?:[a-z0-9-]+\.)+(?:dev|io|me|com|net|org)(?:/[^\s|,;]*)?)",
    re.IGNORECASE,
)


def _extract_text_pdf(path: Path) -> str:
    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text(
                layout=True,
                x_tolerance=2,
                y_tolerance=3,
            )
            if text:
                parts.append(text)
    return "\n".join(parts)


def _extract_text_docx(path: Path) -> str:
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                lines.append(" | ".join(values))
    for relationship in doc.part.rels.values():
        target = str(getattr(relationship, "target_ref", ""))
        if target.startswith(("http://", "https://")):
            lines.append(target)
    return "\n".join(lines)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_text_pdf(path)
    if suffix == ".docx":
        return _extract_text_docx(path)
    if suffix == ".doc":
        raise ValueError("Legacy .doc is not supported — save as .docx or PDF")
    raise ValueError(f"Unsupported file type: {suffix}")


def _normalize_line(line: str) -> str:
    """Collapse letter-spaced text that PDF extraction produces.

    Designers often track out headings, so pdfplumber yields
    'W O R K   E X P E R I E N C E' instead of 'WORK EXPERIENCE'.
    """
    stripped = line.strip()
    tokens = stripped.split()
    if len(tokens) < 3:
        return stripped
    single_char = sum(1 for t in tokens if len(t) == 1)
    if single_char / len(tokens) < 0.6:
        return stripped
    words = [re.sub(r"\s+", "", w) for w in re.split(r"\s{2,}", stripped)]
    return " ".join(w for w in words if w)


def _heading_section(line: str) -> str | None:
    if len(line.split()) > 6 or not line.strip():
        return None
    key = re.sub(r"[^A-Za-z]", "", line).upper()
    if not key or len(key) > 40:
        return None
    for section, names in SECTION_KEYWORDS.items():
        for name in names:
            if key == name or (key.startswith(name) and len(key) <= len(name) + 4):
                return section
    return None


def _split_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {k: [] for k in SECTION_KEYWORDS}
    sections["header"] = []
    current = "header"
    for line in lines:
        section = _heading_section(line)
        if section:
            current = section
            continue
        sections[current].append(line)
    return sections


def _is_contact_line(line: str) -> bool:
    return bool(EMAIL_RE.search(line) or URL_RE.search(line)) or bool(
        PHONE_RE.search(line) and len(line) < 60
    )


def _parse_skills(lines: list[str], text: str, profile_skills: list[str] | None) -> list[str]:
    found: list[str] = []
    if profile_skills:
        lower = text.lower()
        found.extend(s for s in profile_skills if s.lower() in lower)
    for line in lines:
        cleaned = BULLET_PREFIX.sub("", line)
        for piece in re.split(r"[,;|•·/]| {2,}", cleaned):
            piece = piece.strip(" .:-")
            if 1 < len(piece) <= 30 and not piece.isdigit():
                found.append(piece)
    deduped: list[str] = []
    seen: set[str] = set()
    for skill in found:
        key = skill.lower()
        if key not in seen:
            seen.add(key)
            deduped.append(skill)
    return deduped[:40]


def _split_role_company(value: str) -> tuple[str, str]:
    cleaned = value.strip(" |,-–—")
    role, separator, company = cleaned.partition(" at ")
    if separator:
        return role.strip(), company.strip()
    parts = [
        part.strip()
        for part in re.split(r"\s*[|–—]\s*|\s+-\s+", cleaned)
        if part.strip()
    ]
    return (
        parts[0] if parts else "",
        parts[1] if len(parts) > 1 else "",
    )


def _experience_header(
    lines: list[str], index: int
) -> tuple[ExperienceEntry, int] | None:
    line = lines[index].strip()
    if not line or BULLET_PREFIX.match(line):
        return None
    date_match = DATE_RANGE.search(line)
    if date_match:
        role, company = _split_role_company(DATE_RANGE.sub("", line))
        return (
            ExperienceEntry(
                role=role,
                company=company,
                start_date=date_match.group("start").strip(),
                end_date=date_match.group("end").strip().title(),
                bullets=[],
            ),
            1,
        )

    if index + 2 < len(lines):
        company = lines[index + 1].strip()
        date_line = lines[index + 2].strip()
        date_match = DATE_RANGE.search(date_line)
        if (
            date_match
            and not BULLET_PREFIX.match(company)
            and not BULLET_PREFIX.match(date_line)
        ):
            return (
                ExperienceEntry(
                    role=line,
                    company=company,
                    start_date=date_match.group("start").strip(),
                    end_date=date_match.group("end").strip().title(),
                    bullets=[],
                ),
                3,
            )

    if index + 1 < len(lines):
        date_match = DATE_RANGE.search(lines[index + 1])
        if date_match:
            role, company = _split_role_company(line)
            return (
                ExperienceEntry(
                    role=role,
                    company=company,
                    start_date=date_match.group("start").strip(),
                    end_date=date_match.group("end").strip().title(),
                    bullets=[],
                ),
                2,
            )
    return None


def _parse_experience(lines: list[str]) -> list[ExperienceEntry]:
    entries: list[ExperienceEntry] = []
    current: ExperienceEntry | None = None
    index = 0
    while index < len(lines):
        header = _experience_header(lines, index)
        if header:
            current, consumed = header
            entries.append(current)
            index += consumed
            continue

        text = BULLET_PREFIX.sub("", lines[index]).strip()
        if text and current is not None:
            current.bullets.append(text)
        index += 1
    return [entry for entry in entries if entry.role or entry.company][:12]


def _parse_education(lines: list[str]) -> list[EducationEntry]:
    entries: list[EducationEntry] = []
    for line in lines:
        text = BULLET_PREFIX.sub("", line).strip()
        if len(text) < 4:
            continue
        year = ""
        year_match = re.search(r"(19|20)\d{2}", text)
        if year_match:
            year = year_match.group(0)
        without_year = (
            text.replace(year, "").strip(" |,-") if year else text
        )
        parts = [
            part.strip()
            for part in re.split(r"\s*\|\s*", without_year)
            if part.strip()
        ]
        degree = parts[0] if len(parts) > 1 else ""
        institution = parts[1] if len(parts) > 1 else without_year
        entries.append(
            EducationEntry(
                institution=institution,
                degree=degree,
                year=year,
            )
        )
        if len(entries) >= 5:
            break
    return entries


def _parse_projects(lines: list[str]) -> list[ProjectEntry]:
    projects: list[ProjectEntry] = []
    current: ProjectEntry | None = None
    for line in lines:
        text = BULLET_PREFIX.sub("", line).strip()
        if not text:
            continue
        is_bullet = bool(BULLET_PREFIX.match(line))
        if not is_bullet and len(text.split()) <= 14:
            parts = [part.strip() for part in text.split("|") if part.strip()]
            technologies = (
                [
                    tech.strip()
                    for tech in re.split(r"[,;/]", parts[1])
                    if tech.strip()
                ]
                if len(parts) > 1
                else []
            )
            current = ProjectEntry(
                name=parts[0],
                technologies=technologies,
            )
            projects.append(current)
        elif current:
            current.bullets.append(text)
        else:
            current = ProjectEntry(name="Project", description=text)
            projects.append(current)
    return projects[:8]


def _guess_name(header_lines: list[str]) -> str:
    for line in header_lines[:6]:
        clean = line.strip()
        if not clean or _is_contact_line(clean):
            continue
        if 1 <= len(clean.split()) <= 5 and not clean.isupper():
            return clean
        if 1 <= len(clean.split()) <= 5:
            return clean.title()
    return ""


def _build_summary(sections: dict[str, list[str]]) -> str:
    candidates = sections.get("summary") or []
    if not candidates:
        candidates = [
            l for l in sections.get("header", [])
            if len(l.split()) > 8 and not _is_contact_line(l)
        ]
    text = " ".join(l.strip() for l in candidates if not _is_contact_line(l))
    return re.sub(r"\s+", " ", text).strip()[:800]


def _clean_url(value: str) -> str:
    return value.strip().rstrip(".,;:|)")


def parse_resume_text(
    raw: str, profile_skills: list[str] | None = None
) -> ParsedResume:
    lines = [_normalize_line(l) for l in raw.splitlines()]
    lines = [l for l in lines if l.strip()]
    normalized_text = "\n".join(lines)
    sections = _split_sections(lines)

    header_text = "\n".join(sections["header"])
    email = EMAIL_RE.search(header_text)
    phone = PHONE_RE.search(header_text)
    urls = URL_RE.findall(header_text)
    url_list = [_clean_url(u[0] if isinstance(u, tuple) else u) for u in urls]
    linkedin = next((u for u in url_list if "linkedin" in u.lower()), None)
    github = next((u for u in url_list if "github" in u.lower()), None)
    portfolio = next(
        (
            u for u in url_list
            if "linkedin" not in u.lower() and "github" not in u.lower()
        ),
        None,
    )

    experience = _parse_experience(sections["experience"])

    skills = _parse_skills(sections["skills"], normalized_text, profile_skills)

    parsed = ParsedResume(
        name=_guess_name(sections["header"]),
        contact=ContactInfo(
            email=email.group(0) if email else "",
            phone=phone.group(0).strip() if phone else "",
            linkedin=linkedin,
            github=github,
            portfolio=portfolio,
        ),
        summary=_build_summary(sections),
        skills_must_have=skills,
        skills_nice_to_have=[],
        experience=experience,
        projects=_parse_projects(sections["projects"]),
        education=_parse_education(sections["education"]),
        certifications=[
            BULLET_PREFIX.sub("", l).strip() for l in sections["certifications"][:10] if len(l.strip()) > 3
        ],
        achievements=[
            BULLET_PREFIX.sub("", l).strip()
            for l in sections["achievements"][:10]
            if len(l.strip()) > 3
        ],
    )
    if not parsed.name:
        parsed.extraction_warnings.append("Name was not detected")
    if not parsed.experience:
        parsed.extraction_warnings.append("No work experience section was detected")
    if any(not entry.company for entry in parsed.experience):
        parsed.extraction_warnings.append("Some experience entries have no company")
    if parsed.education and any(not entry.degree for entry in parsed.education):
        parsed.extraction_warnings.append("Some education entries have no degree")
    return parsed


def parse_resume_file(
    path: Path, profile_skills: list[str] | None = None
) -> ParsedResume:
    return parse_resume_text(extract_text(path), profile_skills)
