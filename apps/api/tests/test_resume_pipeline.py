from pathlib import Path
from types import SimpleNamespace

from docx import Document

from agents.graphs.tailor_graph import (
    extract_jd_requirements,
    run_tailor_pipeline,
    validate_tailored_resume,
)
from packages.shared.schemas import (
    ContactInfo,
    EducationEntry,
    ExperienceEntry,
    ParsedResume,
)
from services.resume_facts import merge_resume_facts
from services.resume_renderer import render_ats_resume


def _base_resume() -> ParsedResume:
    return ParsedResume(
        name="Ada Example",
        contact=ContactInfo(email="ada@example.com"),
        summary="Backend engineer building reliable services.",
        skills_must_have=["Python", "PostgreSQL", "Docker"],
        experience=[
            ExperienceEntry(
                company="Acme",
                role="Software Engineer",
                start_date="2022",
                end_date="present",
                bullets=[
                    "Built Python APIs serving 20,000 requests per day.",
                    "Reduced database latency by 30% using PostgreSQL indexes.",
                ],
            )
        ],
    )


def test_profile_facts_fill_resume_gaps_without_overwriting_resume_truth():
    profile = SimpleNamespace(
        contact={
            "email": "profile@example.com",
            "phone": "+1 555 0100",
            "linkedin": "linkedin.com/in/ada",
            "github": "github.com/ada",
            "portfolio": "ada.dev",
        },
        locations=["Bangalore", "Remote"],
        education=[
            {"institution": "Example University", "degree": "BSc CS", "year": "2022"}
        ],
        certifications=["AWS Developer"],
        skills_must_have=["Python", "FastAPI"],
        skills_nice_to_have=["Docker"],
        current_role="Software Engineer",
        current_company="Acme",
    )

    merged = merge_resume_facts(profile, _base_resume())

    assert merged.contact.email == "ada@example.com"
    assert merged.contact.phone == "+1 555 0100"
    assert merged.contact.portfolio == "ada.dev"
    assert merged.location == "Bangalore"
    assert merged.education[0].degree == "BSc CS"
    assert merged.certifications == ["AWS Developer"]
    assert "FastAPI" in merged.skills_must_have


def test_jd_extraction_identifies_requirements_and_gaps():
    requirements = extract_jd_requirements(
        "Senior Backend Engineer. 4+ years required. "
        "Must have Python and PostgreSQL. Experience with Kubernetes preferred.",
        "Senior Backend Engineer",
        ["Python", "PostgreSQL", "Docker"],
    )

    assert requirements["experience_years"] == 4
    assert "Python" in requirements["matched_skills"]
    assert "Kubernetes" in requirements["skill_gaps"]


def test_validator_rejects_invented_resume_facts():
    base = _base_resume()
    invented = base.model_copy(deep=True)
    invented.skills_must_have.append("Rust")
    invented.experience[0].company = "Imaginary Corp"

    issues = validate_tailored_resume(base, invented)

    assert any("Invented skills" in issue for issue in issues)
    assert any("Changed or invented companies" in issue for issue in issues)


def test_local_tailor_preserves_facts_and_reports_gaps():
    base = _base_resume()
    tailored_dict, summary = run_tailor_pipeline(
        base.model_dump(),
        "Python backend role using PostgreSQL and Kubernetes. 2+ years.",
        "Backend Engineer",
    )
    tailored = ParsedResume.model_validate(tailored_dict)

    assert tailored.experience[0].company == "Acme"
    assert tailored.experience[0].role == "Software Engineer"
    assert set(tailored.skills_must_have) <= set(base.skills_must_have)
    assert "Kubernetes" in summary


def test_invalid_gemini_output_falls_back_to_local(monkeypatch):
    base = _base_resume()

    def invented_result(*_args, **_kwargs):
        result = base.model_copy(deep=True)
        result.skills_must_have.append("Rust")
        return result

    monkeypatch.setattr(
        "services.gemini_tailor.rewrite_resume_with_gemini",
        invented_result,
    )
    tailored_dict, summary = run_tailor_pipeline(
        base.model_dump(),
        "Python and Rust backend role.",
        "Backend Engineer",
        use_gemini=True,
        gemini_api_key="test-key",
    )

    tailored = ParsedResume.model_validate(tailored_dict)
    assert "Rust" not in tailored.skills_must_have
    assert "local fallback" in summary


def test_renderer_creates_professional_ats_docx(tmp_path: Path):
    resume = _base_resume()
    resume.location = "Bangalore"
    resume.education = [
        EducationEntry(
            institution="Example University", degree="BSc Computer Science", year="2022"
        )
    ]
    resume.certifications = ["AWS Developer"]
    output = tmp_path / "resume.docx"

    render_ats_resume(resume, output)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert output.exists()
    assert "ADA EXAMPLE" in text
    assert "Bangalore" in text
    assert "BSc Computer Science" in text
    assert "AWS Developer" in text
    assert len(document.tables) == 0
